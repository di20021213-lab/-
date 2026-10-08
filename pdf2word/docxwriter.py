"""Запись распознанной структуры в .docx (python-docx)."""

from __future__ import annotations

import io
import math
from statistics import median

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from . import BUILD, __version__, layout, metrics
from . import tables as tbl
from .layout import ColumnGroup, Para, Picture

FONT = "Times New Roman"
ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
TAB = {"left": WD_TAB_ALIGNMENT.LEFT, "right": WD_TAB_ALIGNMENT.RIGHT,
       "center": WD_TAB_ALIGNMENT.CENTER}
CELL_MARGIN_TW = 57           # поля ячеек 0,1 см (у Word по умолчанию 0,19 см) —
                              # меньше риск, что слово не влезет и перенесётся

PAPER = {  # дюймы
    "A4": (8.27, 11.69),
    "A3": (11.69, 16.54),
    "A5": (5.83, 8.27),
    "Letter": (8.5, 11.0),
    "Legal": (8.5, 14.0),
}


class Ctx:
    def __init__(self, dpi: float, body_size: float, font: str = FONT):
        self.dpi = dpi
        self.body_size = body_size
        self.font = font           # шрифт стиля «Обычный»

    def emu(self, px: float) -> Emu:
        return Emu(int(round(px / self.dpi * 914400)))

    def tw(self, px: float) -> int:
        return int(round(px / self.dpi * 1440))

    def pt(self, px: float) -> float:
        return px / self.dpi * 72.0


# --------------------------------------------------------------------------

def _set_fonts(rpr_parent, font: str = FONT, lang: bool = True):
    rpr = rpr_parent.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)
    if not lang:
        return
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rpr.append(el)
    el.set(qn("w:val"), "ru-RU")


def _setup_styles(doc, body_size: float, font: str = FONT):
    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(body_size)
    _set_fonts(normal.element, font)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    try:
        grid = doc.styles["Table Grid"]
        grid.font.name = FONT
    except KeyError:
        pass


def _paper(width_in: float, height_in: float) -> tuple[float, float]:
    """Подгоняет размер листа к стандартному, если он близок."""
    short, long_ = sorted((width_in, height_in))
    for w, h in PAPER.values():
        if abs(short - w) / w < 0.04 and abs(long_ - h) / h < 0.04:
            short, long_ = w, h
            break
    return (long_, short) if width_in > height_in else (short, long_)


def _apply_section(section, pages, ctx: Ctx):
    page = pages[0]
    w_in, h_in = _paper(page.width_px / page.dpi, page.height_px / page.dpi)
    landscape = w_in > h_in
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width = Emu(int(w_in * 914400))
    section.page_height = Emu(int(h_in * 914400))

    def med(vals, default):
        vals = [v for v in vals if v is not None]
        return median(vals) if vals else default

    # левое и правое поля — по заполненным страницам: на последней странице
    # с одной подписью текст обычно не доходит до правого края
    from .pipeline import _all_paras, _text_len
    chars = [sum(_text_len(q) for q in _all_paras(p.elements)) for p in pages]
    full = [p for p, n in zip(pages, chars) if n >= 0.5 * max(chars)] or pages
    lefts = [p.body[0] / p.dpi for p in full]
    rights = [(p.width_px - p.body[2]) / p.dpi for p in full]
    tops = [p.body[1] / p.dpi for p in pages]
    # нижнее поле: по самой заполненной странице
    bottoms = [(p.height_px - p.body[3]) / p.dpi for p in pages]

    def clamp(v, lo=0.3, hi=2.0):
        return max(lo, min(hi, v))

    left = clamp(med(lefts, 1.0) - 0.02)
    right = clamp(med(rights, 0.6) - 0.05)
    top = clamp(min(tops) - 0.05 if tops else 0.8, 0.3, 1.6)
    bottom = clamp(min(bottoms) - 0.05 if bottoms else 0.8, 0.3, 1.6)
    section.left_margin = Emu(int(left * 914400))
    section.right_margin = Emu(int(right * 914400))
    section.top_margin = Emu(int(top * 914400))
    section.bottom_margin = Emu(int(bottom * 914400))
    section.header_distance = Emu(int(min(0.5, top / 2) * 914400))
    section.footer_distance = Emu(int(min(0.5, bottom / 2) * 914400))


# --------------------------------------------------------------------------

def _fill_paragraph(par, p: Para, ctx: Ctx, max_indent_px: float | None = None):
    fmt = par.paragraph_format
    fmt.alignment = ALIGN.get(p.align, WD_ALIGN_PARAGRAPH.LEFT)
    left = p.left
    first = p.first
    if max_indent_px is not None:
        left = min(left, max_indent_px * 0.8)
        first = max(-left, min(first, max_indent_px * 0.6))
    if left > 0:
        fmt.left_indent = ctx.emu(left)
    if abs(first) > 0:
        fmt.first_line_indent = ctx.emu(first)
    if p.space_before > 0:
        fmt.space_before = Pt(round(min(ctx.pt(p.space_before), 200.0), 1))
    if p.spacing and abs(p.spacing - 1.0) > 0.01:
        fmt.line_spacing = p.spacing          # полуторный, двойной интервал
    for pos, kind, *leader in p.tabs:
        fmt.tab_stops.add_tab_stop(ctx.emu(pos), TAB[kind],
                                   WD_TAB_LEADER.DOTS if leader else WD_TAB_LEADER.SPACES)
    size = None if abs(p.size - ctx.body_size) < 0.25 else p.size
    for text, bold, underline, italic, font, script in layout.para_runs(p):
        parts = text.split("\u00ad")
        run = par.add_run(parts[0])
        for part in parts[1:]:
            # мягкий перенос: дефис появится только если слово разорвётся на краю
            run._r.append(OxmlElement("w:softHyphen"))
            t = OxmlElement("w:t")
            t.text = part
            t.set(qn("xml:space"), "preserve")
            run._r.append(t)
        if bold:
            run.bold = True
        if italic:
            run.italic = True
        if underline:
            run.underline = True
        if (font or FONT) != ctx.font:
            _set_fonts(run._r, font or FONT, lang=False)
        if script == "super":
            run.font.superscript = True
        elif script == "sub":
            run.font.subscript = True
        if size:
            run.font.size = Pt(size)
    if size:
        _mark_size(par, size)


def _add_picture(par, pic: Picture, ctx: Ctx, max_width_px: float) -> None:
    """Рисунок в абзаце: размер как на странице, но не шире места под текст."""
    fmt = par.paragraph_format
    width = min(pic.x1 - pic.x0, max_width_px)
    left = 0.0
    if pic.align == "left":
        left = min(pic.left, max(0.0, max_width_px - width))
        if left > 0:
            fmt.left_indent = ctx.emu(left)
    fmt.alignment = ALIGN.get(pic.align, WD_ALIGN_PARAGRAPH.CENTER)
    if pic.space_before > 0:
        fmt.space_before = Pt(round(min(ctx.pt(pic.space_before), 200.0), 1))
    if width < 2:
        return
    try:
        par.add_run().add_picture(io.BytesIO(pic.png), width=ctx.emu(width))
    except Exception:              # испорченная картинка не должна сорвать документ
        return


def _mark_size(par, size: float) -> None:
    """Кегль знака абзаца: от него зависит высота пустой строки и строки таблицы."""
    pPr = par._p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        sect = pPr.find(qn("w:sectPr"))
        if sect is not None:
            sect.addprevious(rPr)
        else:
            pPr.append(rPr)
    for tag in ("w:sz", "w:szCs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rPr.append(el)
        el.set(qn("w:val"), str(int(round(size * 2))))


def _empty_cells_size(table, sizes: list[float], ctx: Ctx) -> None:
    """Пустые ячейки — тем же кеглем, что текст таблицы, иначе строки
    выйдут выше, чем на скане."""
    if not sizes:
        return
    size = median(sizes)
    if abs(size - ctx.body_size) < 0.25:
        return
    for row in table.rows:
        for cell in row.cells:
            pars = cell.paragraphs
            if len(pars) == 1 and not pars[0].text:
                _mark_size(pars[0], size)


def _set_col_widths(table, widths_tw: list[int]):
    tbl_el = table._tbl
    grid = tbl_el.tblGrid
    cols = grid.findall(qn("w:gridCol"))
    for col, w in zip(cols, widths_tw):
        col.set(qn("w:w"), str(max(60, w)))
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            if idx < len(widths_tw):
                cell.width = Emu(int(widths_tw[idx] * 635))
    tblPr = tbl_el.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = _put_tblpr(tblPr, OxmlElement("w:tblW"))
    tblW.set(qn("w:w"), str(sum(widths_tw)))
    tblW.set(qn("w:type"), "dxa")


TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
               "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
               "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]


def _put_tblpr(tblPr, el):
    """Вставляет элемент в w:tblPr в порядке, которого требует схема Word."""
    name = el.tag.split("}")[1]
    old = tblPr.find(qn(f"w:{name}"))
    if old is not None:
        tblPr.remove(old)
    later = TBLPR_ORDER[TBLPR_ORDER.index(name) + 1:]
    for child in tblPr:
        if child.tag.split("}")[1] in later:
            child.addprevious(el)
            return el
    tblPr.append(el)
    return el


def _cell_margins(table, tw: int = CELL_MARGIN_TW):
    tblPr = table._tbl.tblPr
    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is None:
        mar = _put_tblpr(tblPr, OxmlElement("w:tblCellMar"))
    for side in ("left", "right"):
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), str(tw))
        el.set(qn("w:type"), "dxa")


def _cell_borders(cell_obj, right: bool = False):
    """Ячейка без рамки (правая граница — это левая граница самой таблицы)."""
    tcPr = cell_obj._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        # порядок в tcPr: tcW, gridSpan, vMerge, tcBorders, shd, ..., vAlign
        after = [qn("w:tcW"), qn("w:gridSpan"), qn("w:hMerge"), qn("w:vMerge")]
        anchor = None
        for child in tcPr:
            if child.tag in after:
                anchor = child
        if anchor is not None:
            anchor.addnext(borders)
        else:
            tcPr.insert(0, borders)
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        if edge == "right" and right:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "000000")
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)


def _no_borders(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    _put_tblpr(tblPr, borders)


def _cell_widths_merged(cell_obj, width_tw: int):
    cell_obj.width = Emu(int(width_tw * 635))


def _fit_size(p: Para, ctx: Ctx, width_px: float) -> None:
    """Надпись, которая на скане умещалась в ячейке одной строкой, не должна
    переноситься в Word: при нехватке места чуть уменьшаем кегль (до 12 %)."""
    if len(p.lines) != 1 or p.segments is not None or not p.lines[0].words:
        return
    ws = p.lines[0].words
    em = sum(metrics.advance_width(w.text, w.bold) for w in ws) + metrics.SPACE * (len(ws) - 1)
    avail = width_px - 2 * CELL_MARGIN_TW / 1440 * ctx.dpi - p.left - max(0.0, p.first)
    need = em * p.size * ctx.dpi / 72.0
    if need > avail > 0:
        size = math.floor(2 * p.size * avail / need) / 2      # шаг 0,5 пт
        if size >= 0.88 * p.size:
            p.size = size


def _float_picture(par, pic: Picture, ctx: Ctx, max_width_px: float, top_px: float) -> bool:
    """Рисунок, обтекаемый текстом: плавающий, привязан к абзацу par (его
    верх на странице — top_px), стоит у своего края колонки."""
    width = min(pic.x1 - pic.x0, 0.6 * max_width_px)
    if width < 2:
        return False
    run = par.add_run()
    try:
        inline = run.add_picture(io.BytesIO(pic.png), width=ctx.emu(width))._inline
    except Exception:
        par._p.remove(run._r)
        return False
    # рисунок — в начало абзаца, чтобы он вставал рядом с первой строкой
    first = par._p.find(qn("w:r"))
    if first is not None and first is not run._r:
        first.addprevious(run._r)
    gap = str(int(ctx.emu(max(0.0, pic.gap))))
    anchor = OxmlElement("wp:anchor")
    for key, val in (("distT", "0"), ("distB", "0"),
                     ("distL", gap if pic.wrap == "right" else "0"),
                     ("distR", gap if pic.wrap == "left" else "0"),
                     ("simplePos", "0"), ("relativeHeight", "251658240"), ("behindDoc", "0"),
                     ("locked", "0"), ("layoutInCell", "1"), ("allowOverlap", "1")):
        anchor.set(key, val)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    x = min(max(0.0, pic.left), max(0.0, max_width_px - width))
    for tag, rel, off in (("wp:positionH", "column", x), ("wp:positionV", "paragraph", pic.y0 - top_px)):
        pos = OxmlElement(tag)
        pos.set("relativeFrom", rel)
        val = OxmlElement("wp:posOffset")
        val.text = str(int(ctx.emu(off)) if off >= 0 else -int(ctx.emu(-off)))
        pos.append(val)
        anchor.append(pos)
    anchor.append(inline.find(qn("wp:extent")))
    effect = OxmlElement("wp:effectExtent")
    for side in "ltrb":
        effect.set(side, "0")
    anchor.append(effect)
    wrap = OxmlElement("wp:wrapSquare")
    wrap.set("wrapText", "bothSides")
    anchor.append(wrap)
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        el = inline.find(qn(tag))
        if el is not None:
            anchor.append(el)
    inline.getparent().replace(inline, anchor)
    return True


def _write_cell(cell_obj, paras: list[Para], ctx: Ctx, width_px: float, pictures=()):
    first = True
    items = sorted([*paras, *pictures], key=lambda e: e.y0) if pictures else paras
    for p in items:
        if first:
            par = cell_obj.paragraphs[0]
            first = False
        else:
            par = cell_obj.add_paragraph()
        if isinstance(p, Picture):
            p.space_before = 0.0
            p.align, p.left = "center", 0.0
            # с небольшим запасом: иначе Word расширит столбец под рисунок
            _add_picture(par, p, ctx, 0.96 * (width_px - 2 * CELL_MARGIN_TW / 1440 * ctx.dpi))
            continue
        _fit_size(p, ctx, width_px)
        _fill_paragraph(par, p, ctx, max_indent_px=width_px)


def _add_table(doc, t: tbl.Table, ctx: Ctx, max_width_px: float):
    widths_px = [t.xs[i + 1] - t.xs[i] for i in range(t.ncols)]
    total = sum(widths_px)
    scale = min(1.0, max_width_px / total) if total > 0 else 1.0
    widths_tw = [ctx.tw(w * scale) for w in widths_px]
    table = doc.add_table(rows=t.nrows, cols=t.ncols)
    try:
        table.style = doc.styles["Table Grid"]
    except KeyError:
        pass
    table.autofit = False
    where, indent = getattr(t, "place", ("center", 0))
    table.alignment = {"left": WD_TABLE_ALIGNMENT.LEFT, "right": WD_TABLE_ALIGNMENT.RIGHT}.get(
        where, WD_TABLE_ALIGNMENT.CENTER)
    if where == "left" and indent > 0:
        ind = OxmlElement("w:tblInd")
        ind.set(qn("w:w"), str(ctx.tw(min(indent, max(0.0, max_width_px - total)))))
        ind.set(qn("w:type"), "dxa")
        _put_tblpr(table._tbl.tblPr, ind)
    _cell_margins(table)
    _set_col_widths(table, widths_tw)
    if not t.beside_text:
        # строки не ниже, чем на скане: пустые строки бланка остаются
        # местом для записи, высокая шапка — высокой
        for r, row in enumerate(table.rows):
            h = t.ys[r + 1] - t.ys[r]
            if h > 0:
                row.height = ctx.emu(0.97 * h)
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    for c in t.cells:
        a = table.cell(c.r0, c.c0)
        if c.r1 - c.r0 > 1 or c.c1 - c.c0 > 1:
            a = a.merge(table.cell(c.r1 - 1, c.c1 - 1))
            _cell_widths_merged(a, sum(widths_tw[c.c0:c.c1]))
    _empty_cells_size(table, [p.size for c in t.cells for p in c.paragraphs], ctx)
    for c in t.cells:
        cell_obj = table.cell(c.r0, c.c0)
        _write_cell(cell_obj, c.paragraphs, ctx, (c.x1 - c.x0) * scale, c.pictures)
        if c.borderless:
            _cell_borders(cell_obj, right=True)
        if c.valign:
            cell_obj.vertical_alignment = {
                "top": WD_CELL_VERTICAL_ALIGNMENT.TOP, "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
                "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER}[c.valign]
            continue
        if c.paragraphs:
            top_gap = c.paragraphs[0].y0 - c.y0
            bottom_gap = c.y1 - c.paragraphs[-1].y1
            h = c.y1 - c.y0
            if h > 0 and top_gap > 0.2 * h and abs(top_gap - bottom_gap) < 0.25 * h:
                cell_obj.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            elif h > 0 and top_gap > 0.4 * h and bottom_gap < 0.2 * h:
                cell_obj.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    return table


def _add_columns(doc, g: ColumnGroup, ctx: Ctx, max_width_px: float):
    widths_px = [b - a for a, b in g.bounds]
    # колонка не уже своей самой длинной строки вместе с полями ячейки, иначе
    # Word перенесёт строку; недостающее место берём у соседней колонки
    pad = 2 * CELL_MARGIN_TW / 1440 * ctx.dpi + 0.03 * ctx.dpi
    need = [max((ln.x1 for p in paras for ln in p.lines), default=a) - a + pad if paras else 0.0
            for (a, _), paras in zip(g.bounds, g.cells)]
    for k in range(len(widths_px)):
        short = need[k] - widths_px[k]
        for j in (k + 1, k - 1):
            if short > 0 and 0 <= j < len(widths_px) and widths_px[j] > need[j]:
                d = min(widths_px[j] - need[j], short)
                widths_px[j] -= d
                widths_px[k] += d
                short -= d
    total = sum(widths_px)
    if 0 < total < max_width_px:
        # колонки растягиваем до полной ширины текста (запас справа)
        widths_px[-1] += max_width_px - total
        total = max_width_px
    scale = min(1.0, max_width_px / total) if total > 0 else 1.0
    widths_tw = [ctx.tw(w * scale) for w in widths_px]
    table = doc.add_table(rows=1, cols=len(widths_tw))
    table.autofit = False
    _no_borders(table)
    _cell_margins(table)
    _set_col_widths(table, widths_tw)
    _empty_cells_size(table, [p.size for paras in g.cells for p in paras], ctx)
    for k, paras in enumerate(g.cells):
        _write_cell(table.cell(0, k), paras, ctx, widths_px[k] * scale)
    return table


def _spacer(doc, ctx: Ctx, gap_px: float):
    par = doc.add_paragraph()
    fmt = par.paragraph_format
    h = ctx.pt(gap_px)
    if h >= 4:
        fmt.line_spacing = Pt(min(h, 400))
        from docx.enum.text import WD_LINE_SPACING
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    else:
        fmt.line_spacing = Pt(1)
        from docx.enum.text import WD_LINE_SPACING
        fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    return par


def _main_font(flow) -> str:
    """Шрифт, которым набрана большая часть текста (у сканов — Times New Roman)."""
    count: dict[str, int] = {}

    def add(paras):
        for p in paras:
            for ln in p.lines:
                for w in ln.words:
                    f = w.font or FONT
                    count[f] = count.get(f, 0) + len(w.text)

    for el in flow:
        if isinstance(el, Para):
            add([el])
        elif isinstance(el, ColumnGroup):
            for paras in el.cells:
                add(paras)
        elif isinstance(el, tbl.Table):
            for c in el.cells:
                add(c.paragraphs)
    return max(count, key=count.get) if count else FONT


# --------------------------------------------------------------------------

def _auto_hyphenation(doc) -> None:
    """Автоматический перенос слов (по правилам языка текста)."""
    settings = doc.settings.element
    if settings.find(qn("w:autoHyphenation")) is not None:
        return
    el = OxmlElement("w:autoHyphenation")
    # место в w:settings задано схемой: сразу после w:defaultTabStop
    anchor = settings.find(qn("w:defaultTabStop"))
    if anchor is not None:
        anchor.addnext(el)
        return
    for tag in ("w:characterSpacingControl", "w:compat", "w:rsids", "w:themeFontLang"):
        nxt = settings.find(qn(tag))
        if nxt is not None:
            nxt.addprevious(el)
            return
    settings.append(el)


def write(pages, flow, out_path: str, title: str = "", hyphenate: bool = False):
    from .pipeline import PageBreak, SectionBreak

    doc = Document()
    if not pages:
        doc.save(out_path)
        return
    sizes = [p.body_size for p in pages]
    body_size = median(sizes)
    dpi = median([p.dpi for p in pages])
    font = _main_font(flow)
    ctx = Ctx(dpi, body_size, font)
    _setup_styles(doc, body_size, font)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")
    doc.core_properties.title = title
    doc.core_properties.author = "PDF в Word"
    # какой версией сделан файл (видно в «Файл → Сведения» у Word)
    doc.core_properties.version = __version__
    doc.core_properties.comments = f"PDF в Word {__version__}" + (f", сборка {BUILD}" if BUILD else "")
    if hyphenate:
        _auto_hyphenation(doc)

    # группы страниц с одинаковой ориентацией → разделы
    groups = [[pages[0]]]
    for p in pages[1:]:
        if (p.width_px > p.height_px) == (groups[-1][-1].width_px > groups[-1][-1].height_px):
            groups[-1].append(p)
        else:
            groups.append([p])
    section = doc.sections[0]
    _apply_section(section, groups[0], ctx)
    group_iter = iter(groups[1:])

    def text_width_px(sec):
        w = sec.page_width - sec.left_margin - sec.right_margin
        return w / 914400 * dpi

    max_w = text_width_px(section)
    body = doc.element.body
    # убираем пустой абзац шаблона, если он есть
    for par in list(doc.paragraphs):
        if not par.text.strip():
            body.remove(par._p)

    prev_kind = None
    new_page = False
    floating = None                     # рисунок с обтеканием ждёт своего абзаца
    for el in flow:
        if floating is not None and not isinstance(el, Para) and \
                not (isinstance(el, Picture) and el.wrap):
            _add_picture(doc.add_paragraph(), floating, ctx, max_w)
            floating = None
        if isinstance(el, SectionBreak):
            grp = next(group_iter, None)
            section = doc.add_section(WD_SECTION.NEW_PAGE)
            if grp:
                _apply_section(section, grp, ctx)
            max_w = text_width_px(section)
            prev_kind = None
            new_page = False
            continue
        if isinstance(el, PageBreak):
            new_page = prev_kind is not None
            continue
        if isinstance(el, Para):
            par = doc.add_paragraph()
            _fill_paragraph(par, el, ctx, max_indent_px=max_w)
            par.paragraph_format.page_break_before = new_page or None
            if floating is not None:
                if not _float_picture(par, floating, ctx, max_w, el.y0):
                    _add_picture(doc.add_paragraph(), floating, ctx, max_w)
                floating = None
            prev_kind = "para"
        elif isinstance(el, Picture) and el.wrap:
            if floating is not None:
                _add_picture(doc.add_paragraph(), floating, ctx, max_w)
            floating = el
            continue
        elif isinstance(el, Picture):
            par = doc.add_paragraph()
            _add_picture(par, el, ctx, max_w)
            par.paragraph_format.page_break_before = new_page or None
            prev_kind = "para"
        elif isinstance(el, (tbl.Table, ColumnGroup)):
            gap = getattr(el, "space_before", 0.0) or 0.0
            if new_page:
                _spacer(doc, ctx, 1.0).paragraph_format.page_break_before = True
            elif prev_kind in ("table", "cols"):
                _spacer(doc, ctx, max(gap, 2.0))
            elif gap > 0.6 * ctx.body_size * dpi / 72:
                _spacer(doc, ctx, gap)
            if isinstance(el, tbl.Table):
                _add_table(doc, el, ctx, max_w)
                prev_kind = "table"
            else:
                _add_columns(doc, el, ctx, max_w)
                prev_kind = "cols"
        new_page = False
    if floating is not None:
        _add_picture(doc.add_paragraph(), floating, ctx, max_w)
        prev_kind = "para"
    # Word требует абзац после таблицы в конце документа
    if prev_kind in ("table", "cols"):
        _spacer(doc, ctx, 2.0)
    doc.save(out_path)
