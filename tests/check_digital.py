"""Проверяет .docx, полученный из электронного PDF (tests/make_digital.py):
текст точный, перенос по слогам склеен, колонтитулы убраны, есть таблица,
рисунок, жирный и курсив.

    python tests/check_digital.py digital.docx digital.txt
"""

from __future__ import annotations

import difflib
import sys

from docx import Document
from docx.oxml.ns import qn


def paragraph_text(p) -> str:
    out = []
    for el in p.iter():
        if el.tag == qn("w:t"):
            out.append(el.text or "")
        elif el.tag == qn("w:tab"):
            out.append(" ")
    return "".join(out)          # мягкий перенос (w:softHyphen) не даёт текста


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError, OSError):
            pass
    docx_path, truth_path = sys.argv[1], sys.argv[2]
    doc = Document(docx_path)
    body = doc.element.body
    paras = list(body.iter(qn("w:p")))
    got = " ".join(paragraph_text(p) for p in paras).split()
    truth = open(truth_path, encoding="utf-8").read().split()
    sm = difflib.SequenceMatcher(None, truth, got, autojunk=False)
    matched = sum(b.size for b in sm.get_matching_blocks())
    recall = matched / max(1, len(truth))
    print(f"слов в эталоне: {len(truth)}, в документе: {len(got)}, совпало: {matched} ({recall:.1%})")
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != "equal":
            print(f"  {tag}: {' '.join(truth[i1:i2])!r} -> {' '.join(got[j1:j2])!r}")

    text = " ".join(got)
    checks = {
        "текст совпадает": recall >= 0.98 and len(got) <= len(truth) + 2,
        "перенос склеен": "вычислительной" in got,
        "колонтитулы убраны": "пособие" not in text,
        "таблица 3 столбца": bool(doc.tables) and len(doc.tables[0].columns) == 3,
        "рисунок": len(list(body.iter(qn("w:drawing")))) >= 1,
        "жирный": any(r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:b")) is not None
                      and "Переменная" in "".join(t.text or "" for t in r.iter(qn("w:t")))
                      for r in body.iter(qn("w:r"))),
        "курсив": any(r.find(qn("w:rPr")) is not None and r.find(qn("w:rPr")).find(qn("w:i")) is not None
                      and "памяти" in "".join(t.text or "" for t in r.iter(qn("w:t")))
                      for r in body.iter(qn("w:r"))),
    }
    for name, ok in checks.items():
        print(f"  {'+' if ok else '-'} {name}")
    ok = all(checks.values())
    print("ИТОГ:", "OK" if ok else "ОШИБКА")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
