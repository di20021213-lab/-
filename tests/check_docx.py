"""Проверяет распознанный .docx по эталонным словам тестового скана.

    python tests/check_docx.py result.docx test_scan.txt [мин_доля]
"""

from __future__ import annotations

import difflib
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn


def words_of(path: str) -> list[str]:
    doc = Document(path)
    out: list[str] = []
    for el in doc.element.body.iterchildren():
        if el.tag in (qn("w:p"), qn("w:tbl")):
            for t in el.iter(qn("w:t")):
                out += (t.text or "").split()
    return out


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError, OSError):
            pass
    docx_path, truth_path = sys.argv[1], sys.argv[2]
    need = float(sys.argv[3]) if len(sys.argv) > 3 else 0.9
    truth = open(truth_path, encoding="utf-8").read().split()
    got = words_of(docx_path)
    sm = difflib.SequenceMatcher(None, truth, got, autojunk=False)
    matched = sum(b.size for b in sm.get_matching_blocks())
    recall = matched / max(1, len(truth))
    doc = Document(docx_path)
    tables = doc.tables
    orients = [s.orientation for s in doc.sections]
    print(f"слов в эталоне: {len(truth)}, распознано: {len(got)}, совпало: {matched} ({recall:.1%})")
    print(f"таблиц: {len(tables)}, колонок в первой: {len(tables[0].columns) if tables else 0}")
    print("разделы:", ["альбомный" if o == WD_ORIENT.LANDSCAPE else "книжный" for o in orients])
    for tag, i1, i2, j1, j2 in sm.get_opcodes():
        if tag != "equal":
            print(f"  {tag}: {' '.join(truth[i1:i2])!r} -> {' '.join(got[j1:j2])!r}")
    ok = recall >= need and tables and len(tables[0].columns) == 3 and \
        WD_ORIENT.LANDSCAPE in orients and WD_ORIENT.PORTRAIT in orients
    print("ИТОГ:", "OK" if ok else "ОШИБКА")
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
